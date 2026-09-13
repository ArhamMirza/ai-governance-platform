"""Reading text out of the file types this platform accepts as evidence.

Every case here is checked for the same thing twice over: did it get the
text, and -- when it could not -- did it say why. The second half is the
point. A file that indexes as silently empty is indistinguishable from a
file with nothing relevant in it, and an assistant that cannot tell those
apart will report "the record doesn't cover that" about a signed approval
sitting in the folder.
"""

import io

import pytest
from PIL import Image, ImageDraw

from app.core import ocr
from app.models.document_chunk import ChunkSource, IndexStatus
from app.services.document_extraction import extract
from tests.conftest import fake_image


def docx_bytes(paragraphs: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for style, text in paragraphs:
        document.add_paragraph(text, style=style) if style else document.add_paragraph(text)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row_cells, values in zip(added.rows, table, strict=False):
            for cell, value in zip(row_cells.cells, values, strict=False):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(sheets: dict[str, list[list[str]]]) -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def image_only_pdf(lines: list[str]) -> bytes:
    """A scan: a PDF page that is a picture of text, with no text layer."""
    image = Image.new("RGB", (1200, 400), "white")
    draw = ImageDraw.Draw(image)
    for n, line in enumerate(lines):
        draw.text((30, 40 + n * 60), line, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PDF")
    return buffer.getvalue()


# ---------- formats that carry their text ----------


def test_a_text_file_is_read() -> None:
    result = extract("notes.txt", b"The retention period is 24 months.")

    assert result.status is IndexStatus.INDEXED
    assert "24 months" in result.segments[0].text


def test_a_csv_row_keeps_its_column_names() -> None:
    """A retrieved row has to say what its values mean; bare cells are
    unreadable on their own and unciteable."""
    result = extract("controls.csv", b"Control,Status\nA.5.23,Implemented\n")

    assert "Control: A.5.23" in result.segments[0].text
    assert "Status: Implemented" in result.segments[0].text


def test_a_spreadsheet_is_split_by_sheet() -> None:
    result = extract(
        "register.xlsx",
        xlsx_bytes({"Risks": [["Risk", "Owner"], ["Bias", "Ada"]], "Controls": [["A.5.23"]]}),
    )

    assert [s.section for s in result.segments] == ["Sheet: Risks", "Sheet: Controls"]
    assert "Bias" in result.segments[0].text


def test_a_word_document_is_split_by_heading() -> None:
    """Headings become chunk boundaries and the label a citation quotes."""
    result = extract(
        "policy.docx",
        docx_bytes(
            [
                ("Heading 1", "Data Protection"),
                (None, "Personal data is retained for 24 months."),
                ("Heading 1", "Access Control"),
                (None, "Access is reviewed quarterly."),
            ]
        ),
    )

    assert [s.section for s in result.segments] == ["Data Protection", "Access Control"]
    assert "24 months" in result.segments[0].text


def test_a_table_in_a_word_document_is_not_lost() -> None:
    result = extract(
        "policy.docx",
        docx_bytes([(None, "See the table.")], table=[["Control", "Owner"], ["A.5.23", "Ada"]]),
    )

    tables = [s for s in result.segments if s.section == "Table 1"]
    assert tables and "A.5.23" in tables[0].text


# ---------- pictures of text ----------


def test_a_photograph_of_a_document_is_read_rather_than_skipped(monkeypatch) -> None:
    """The case that motivates OCR at all: evidence uploaded as a picture."""
    result = extract("signed-approval.png", fake_image("Approved by the Data Owner"))

    assert result.status is IndexStatus.INDEXED
    assert result.segments[0].source is ChunkSource.OCR
    assert "Data Owner" in result.segments[0].text
    # The reader is told the text came from a picture.
    assert "OCR" in (result.detail or "")


def test_a_scanned_pdf_is_read_through_its_page_images(monkeypatch) -> None:
    """A real PDF whose page has no text layer. Without this it extracts as
    empty -- and scanned approvals are among the commonest evidence there is.

    pypdf and the PDF itself are real here; only the OCR engine is stubbed,
    so the page-image path is genuinely exercised.
    """
    monkeypatch.setattr(ocr, "read_image", lambda data: "Retention period: 24 months")

    result = extract("scan.pdf", image_only_pdf(["Retention period: 24 months"]))

    assert result.status is IndexStatus.INDEXED
    assert result.segments[0].source is ChunkSource.OCR
    assert result.segments[0].section == "page 1"
    assert "24 months" in result.segments[0].text
    assert "scan" in (result.detail or "").lower()


def test_an_image_with_no_legible_text_says_so(monkeypatch) -> None:
    monkeypatch.setattr(ocr, "read_image", lambda data: "")

    result = extract("wall.jpg", b"anything")

    assert result.status is IndexStatus.EMPTY
    assert "legible" in result.detail


def test_images_are_reported_as_unread_when_ocr_is_unavailable(monkeypatch) -> None:
    """Degrading silently here would be the worst option: the file looks
    indexed and answers nothing."""
    monkeypatch.setattr(ocr, "available", lambda: False)

    result = extract("photo.png", fake_image("anything"))

    assert result.status is IndexStatus.EMPTY
    assert "OCR is unavailable" in result.detail


# ---------- files that cannot be read ----------


@pytest.mark.parametrize(
    "filename, expected_in_detail",
    [
        ("old.doc", "Word 97-2003"),
        ("old.xls", "Excel 97-2003"),
    ],
)
def test_a_legacy_office_file_names_the_format_and_the_fix(
    filename: str, expected_in_detail: str
) -> None:
    result = extract(filename, b"\xd0\xcf\x11\xe0binary")

    assert result.status is IndexStatus.UNREADABLE
    assert expected_in_detail in result.detail
    assert "Re-save" in result.detail


def test_a_corrupt_file_is_recorded_not_raised() -> None:
    """A malformed upload is a fact about the document, not a server error."""
    result = extract("report.pdf", b"this is not a pdf at all")

    assert result.status is IndexStatus.UNREADABLE
    assert result.detail


def test_an_empty_file_is_empty_rather_than_broken() -> None:
    assert extract("blank.txt", b"   \n  ").status is IndexStatus.EMPTY


def test_an_unknown_extension_is_refused_by_name() -> None:
    result = extract("archive.zip", b"PK\x03\x04")

    assert result.status is IndexStatus.UNREADABLE
    assert ".zip" in result.detail

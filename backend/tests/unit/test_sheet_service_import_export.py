import io
from types import SimpleNamespace

import openpyxl
import pytest
from sqlalchemy.orm import Session

from app.core.exceptions import ValidationError
from app.schemas.sheet import ColumnDef
from app.services import sheet_service


def make_upload_file(filename: str, content: bytes):
    return SimpleNamespace(filename=filename, file=io.BytesIO(content))


def build_xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_import_csv_creates_columns_and_rows(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(db_session, "RiskAssessment", "Risk Items", [])
    csv_content = b"Category,Likelihood\nBias,3\nPrivacy,5\n"
    upload = make_upload_file("risks.csv", csv_content)

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert [c["key"] for c in outcome.sheet.columns] == ["category", "likelihood"]
    rows = sheet_service.list_rows(db_session, sheet.id)
    assert [r.data for r in rows] == [
        {"category": "Bias", "likelihood": "3"},
        {"category": "Privacy", "likelihood": "5"},
    ]


def test_import_xlsx_creates_columns_and_rows(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(db_session, "RiskAssessment", "Risk Items", [])
    xlsx_content = build_xlsx_bytes(
        ["Category", "Likelihood"], [["Bias", 3], ["Privacy", 5]]
    )
    upload = make_upload_file("risks.xlsx", xlsx_content)

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert [c["key"] for c in outcome.sheet.columns] == ["category", "likelihood"]
    rows = sheet_service.list_rows(db_session, sheet.id)
    assert [r.data for r in rows] == [
        {"category": "Bias", "likelihood": "3"},
        {"category": "Privacy", "likelihood": "5"},
    ]


def test_import_leaves_an_unmatched_row_alone(db_session: Session) -> None:
    """A row the file doesn't mention -- no id, no name match -- survives
    the import. Replacing it required an explicit delete before; this
    session's audit found that "replace everything" is what silently
    misattributed an archived AI system's life cycle to an unrelated row
    that happened to reuse its id."""
    sheet = sheet_service.create_sheet(
        db_session, "RiskAssessment", "Risk Items", [ColumnDef(key="old", label="Old")]
    )
    kept = sheet_service.add_row(db_session, sheet.id, {"old": "stale"})
    upload = make_upload_file("risks.csv", b"Category\nBias\n")

    sheet_service.import_upload(db_session, sheet.id, upload)

    rows = sheet_service.list_rows(db_session, sheet.id)
    assert len(rows) == 2
    assert {r.id: r.data for r in rows}[kept.id] == {"old": "stale"}
    assert {"category": "Bias"} in [r.data for r in rows]


def test_import_rejects_empty_csv_file(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(db_session, "RiskAssessment", "Risk Items", [])
    upload = make_upload_file("risks.csv", b"")

    with pytest.raises(ValidationError):
        sheet_service.import_upload(db_session, sheet.id, upload)


def test_import_rejects_unsupported_file_type(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(db_session, "RiskAssessment", "Risk Items", [])
    upload = make_upload_file("risks.txt", b"not a real sheet")

    with pytest.raises(ValidationError):
        sheet_service.import_upload(db_session, sheet.id, upload)


def test_import_deduplicates_repeated_headers(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(db_session, "RiskAssessment", "Risk Items", [])
    upload = make_upload_file("risks.csv", b"Note,Note\nfirst,second\n")

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert [c["key"] for c in outcome.sheet.columns] == ["note", "note_1"]


def test_import_keeps_a_server_owned_column_the_file_omits(db_session: Session) -> None:
    """An uploaded file has no reason to carry lifecycle_stage, but the stage
    itself still exists -- dropping the column would leave it with nowhere to
    show and let a re-upload quietly undo the reserved-column guard."""
    sheet = sheet_service.create_sheet(
        db_session,
        "AISystemInventory",
        "AI Systems",
        [ColumnDef(key="name", label="Name"), ColumnDef(key="lifecycle_stage", label="Stage")],
    )
    upload = make_upload_file("systems.csv", b"Name\nResume Screener\n")

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert [c["key"] for c in outcome.sheet.columns] == ["name", "lifecycle_stage"]
    # The label the app gave it survives too, rather than being regenerated.
    assert outcome.sheet.columns[1]["label"] == "Stage"


def test_import_keeps_the_data_inventory_review_column(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(
        db_session,
        "DataInventory",
        "Data Inventory",
        [
            ColumnDef(key="dataset_name", label="Dataset Name"),
            ColumnDef(key="approval_status", label="Approval Status"),
        ],
    )
    upload = make_upload_file("datasets.csv", b"Dataset Name\nApplicant CVs\n")

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert "approval_status" in {c["key"] for c in outcome.sheet.columns}


def test_import_does_not_duplicate_a_server_owned_column_the_file_includes(
    db_session: Session,
) -> None:
    sheet = sheet_service.create_sheet(
        db_session,
        "AISystemInventory",
        "AI Systems",
        [ColumnDef(key="lifecycle_stage", label="Stage")],
    )
    upload = make_upload_file("systems.csv", b"Name,Lifecycle Stage\nResume Screener,design\n")

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert [c["key"] for c in outcome.sheet.columns] == ["name", "lifecycle_stage"]


def test_import_does_not_invent_reserved_columns_the_sheet_never_had(
    db_session: Session,
) -> None:
    """Only a reserved column the sheet actually carries is preserved -- a risk
    sheet with no scoring columns doesn't grow them on import."""
    sheet = sheet_service.create_sheet(
        db_session, "RiskAssessment", "Risk Items", [ColumnDef(key="old", label="Old")]
    )
    upload = make_upload_file("risks.csv", b"Category\nBias\n")

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert [c["key"] for c in outcome.sheet.columns] == ["category"]


def test_import_keeps_the_risk_scoring_columns(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(
        db_session,
        "RiskAssessment",
        "Risk Items",
        [
            ColumnDef(key="category", label="Category"),
            ColumnDef(key="likelihood", label="Likelihood"),
            ColumnDef(key="impact", label="Impact"),
        ],
    )
    upload = make_upload_file("risks.csv", b"Category\nBias\n")

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert [c["key"] for c in outcome.sheet.columns] == ["category", "likelihood", "impact"]


def test_export_xlsx_round_trips_rows(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(
        db_session,
        "RiskAssessment",
        "Risk Items",
        [
            ColumnDef(key="category", label="Category"),
            ColumnDef(key="likelihood", label="Likelihood"),
        ],
    )
    row = sheet_service.add_row(db_session, sheet.id, {"category": "Bias", "likelihood": 3})

    exported = sheet_service.export_xlsx(db_session, sheet.id)

    workbook = openpyxl.load_workbook(io.BytesIO(exported))
    worksheet = workbook.active
    values = list(worksheet.iter_rows(values_only=True))
    assert values == [("ID", "Category", "Likelihood"), (row.id, "Bias", 3)]


# ---------- re-import matches rows back to themselves ----------
#
# This is the fix: sheet_service.import_upload used to delete every row and
# recreate all of them, and SQLite reuses a freed rowid -- so a life cycle,
# an assessment, or a dataset's review could silently end up attached to
# whatever unrelated row happened to reuse its old id on the next re-upload.
# These match a re-imported row back to the one it actually is, by the
# file's own ID column (round-tripped through export_xlsx) or, failing
# that, a configured natural-key column.


def test_reimport_by_id_updates_the_matching_row_in_place(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(
        db_session, "RiskAssessment", "Risk Items", [ColumnDef(key="category", label="Category")]
    )
    row = sheet_service.add_row(db_session, sheet.id, {"category": "Bias"})

    exported = sheet_service.export_xlsx(db_session, sheet.id)
    workbook = openpyxl.load_workbook(io.BytesIO(exported))
    workbook.active["B2"] = "Privacy"  # edit the Category cell in place
    buffer = io.BytesIO()
    workbook.save(buffer)
    upload = make_upload_file("edited.xlsx", buffer.getvalue())

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    rows = sheet_service.list_rows(db_session, sheet.id)
    assert [r.id for r in rows] == [row.id]
    assert rows[0].data == {"category": "Privacy"}
    assert outcome.changed_row_ids == {row.id}


def test_reimport_of_an_unchanged_row_reports_no_change(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(
        db_session, "RiskAssessment", "Risk Items", [ColumnDef(key="category", label="Category")]
    )
    sheet_service.add_row(db_session, sheet.id, {"category": "Bias"})
    exported = sheet_service.export_xlsx(db_session, sheet.id)
    upload = make_upload_file("same.xlsx", exported)

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert outcome.changed_row_ids == set()


def test_the_id_column_never_becomes_a_real_column(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(
        db_session, "RiskAssessment", "Risk Items", [ColumnDef(key="category", label="Category")]
    )
    sheet_service.add_row(db_session, sheet.id, {"category": "Bias"})
    exported = sheet_service.export_xlsx(db_session, sheet.id)
    upload = make_upload_file("same.xlsx", exported)

    outcome = sheet_service.import_upload(db_session, sheet.id, upload)

    assert "id" not in {c["key"] for c in outcome.sheet.columns}
    rows = sheet_service.list_rows(db_session, sheet.id)
    assert "id" not in rows[0].data


def test_an_id_that_does_not_resolve_creates_a_new_row_instead_of_failing(
    db_session: Session,
) -> None:
    sheet = sheet_service.create_sheet(
        db_session, "RiskAssessment", "Risk Items", [ColumnDef(key="category", label="Category")]
    )
    upload = make_upload_file("hand-built.csv", b"ID,Category\n999,Bias\n")

    sheet_service.import_upload(db_session, sheet.id, upload)

    rows = sheet_service.list_rows(db_session, sheet.id)
    assert [r.data for r in rows] == [{"category": "Bias"}]
    assert rows[0].id != 999


def test_reimport_falls_back_to_the_natural_key_when_the_file_has_no_id_column(
    db_session: Session,
) -> None:
    sheet = sheet_service.create_sheet(
        db_session,
        "AISystemInventory",
        "AI Systems",
        [ColumnDef(key="name", label="Name"), ColumnDef(key="purpose", label="Purpose")],
    )
    row = sheet_service.add_row(db_session, sheet.id, {"name": "Resume Screener", "purpose": ""})
    upload = make_upload_file(
        "systems.csv", b"Name,Purpose\nResume Screener,Rank applicants\n"
    )

    outcome = sheet_service.import_upload(
        db_session, sheet.id, upload, natural_key_column="name"
    )

    rows = sheet_service.list_rows(db_session, sheet.id)
    assert [r.id for r in rows] == [row.id]
    assert rows[0].data["purpose"] == "Rank applicants"
    assert outcome.changed_row_ids == {row.id}


def test_reimport_name_match_is_case_insensitive_and_trimmed(db_session: Session) -> None:
    sheet = sheet_service.create_sheet(
        db_session, "AISystemInventory", "AI Systems", [ColumnDef(key="name", label="Name")]
    )
    row = sheet_service.add_row(db_session, sheet.id, {"name": "Resume Screener"})
    upload = make_upload_file("systems.csv", b"Name\n  resume screener  \n")

    sheet_service.import_upload(db_session, sheet.id, upload, natural_key_column="name")

    rows = sheet_service.list_rows(db_session, sheet.id)
    assert [r.id for r in rows] == [row.id]


def test_a_name_not_in_the_file_is_left_alone_and_a_new_name_is_created(
    db_session: Session,
) -> None:
    sheet = sheet_service.create_sheet(
        db_session, "AISystemInventory", "AI Systems", [ColumnDef(key="name", label="Name")]
    )
    existing = sheet_service.add_row(db_session, sheet.id, {"name": "Resume Screener"})
    upload = make_upload_file("systems.csv", b"Name\nFraud Detector\n")

    sheet_service.import_upload(db_session, sheet.id, upload, natural_key_column="name")

    rows = {r.data["name"] for r in sheet_service.list_rows(db_session, sheet.id)}
    assert rows == {"Resume Screener", "Fraud Detector"}
    still_there = sheet_service.list_rows(db_session, sheet.id)
    assert existing.id in {r.id for r in still_there}


def test_without_a_natural_key_configured_an_unmatched_file_row_is_never_guessed_at(
    db_session: Session,
) -> None:
    """Risk items have nothing as reliably unique as a name -- category and
    description both repeat -- so only the ID column may match a row back
    to itself; a file with none just creates fresh rows."""
    sheet = sheet_service.create_sheet(
        db_session, "RiskAssessment", "Risk Items", [ColumnDef(key="category", label="Category")]
    )
    existing = sheet_service.add_row(db_session, sheet.id, {"category": "Bias"})
    upload = make_upload_file("risks.csv", b"Category\nBias\n")

    sheet_service.import_upload(db_session, sheet.id, upload)

    rows = sheet_service.list_rows(db_session, sheet.id)
    assert len(rows) == 2
    assert existing.id in {r.id for r in rows}

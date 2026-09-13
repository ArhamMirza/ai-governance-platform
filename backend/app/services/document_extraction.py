"""Getting readable text out of an uploaded file.

Every path here ends in one of two honest outcomes: text with a structural
label attached to it, or a recorded reason why there is none. What it never
does is return empty text without saying why -- that is the case that makes
an assistant claim the record is silent about something the record actually
documents.

Structure is preserved as far as each format offers it (page, sheet, heading),
because ``app.core.chunking`` uses those boundaries and because a citation is
far more useful saying "page 4" than "somewhere in this PDF".
"""

import csv
import io
from dataclasses import dataclass
from pathlib import Path

from app.core import ocr
from app.core.chunking import Segment
from app.models.document_chunk import ChunkSource, IndexStatus

# A PDF page with less than this much text is treated as a scan of a page
# rather than a page, and its images are read instead.
MIN_PDF_PAGE_CHARS = 40
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
# Formats allowed as evidence that nothing here can parse. Named explicitly
# so the reason given to the user is specific rather than "failed".
LEGACY_OFFICE = {".doc": "Word 97-2003", ".xls": "Excel 97-2003"}


@dataclass
class Extraction:
    segments: list[Segment]
    status: IndexStatus
    detail: str | None = None

    @classmethod
    def failed(cls, detail: str) -> "Extraction":
        return cls(segments=[], status=IndexStatus.UNREADABLE, detail=detail)

    @classmethod
    def empty(cls, detail: str) -> "Extraction":
        return cls(segments=[], status=IndexStatus.EMPTY, detail=detail)

    @classmethod
    def of(cls, segments: list[Segment]) -> "Extraction":
        useful = [s for s in segments if s.text.strip()]
        if not useful:
            return cls.empty("No text found in this file.")
        return cls(segments=useful, status=IndexStatus.INDEXED)


def _from_text(data: bytes) -> Extraction:
    return Extraction.of([Segment(text=data.decode("utf-8", errors="replace"))])


def _from_csv(data: bytes) -> Extraction:
    text = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return Extraction.empty("The file has no rows.")
    # Each row carries its own header labels, so a retrieved row still says
    # what its values mean instead of being a bare list of cells.
    header, *body = rows
    lines = [
        "; ".join(f"{h}: {v}" for h, v in zip(header, row, strict=False) if v.strip())
        for row in body
    ]
    return Extraction.of([Segment(text="\n\n".join(line for line in lines if line))])


def _from_xlsx(data: bytes) -> Extraction:
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return Extraction.failed(f"Could not open the workbook ({type(exc).__name__}).")

    segments = []
    for sheet in book.worksheets:
        lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
        if lines:
            segments.append(Segment(text="\n\n".join(lines), section=f"Sheet: {sheet.title}"))
    book.close()
    return Extraction.of(segments)


def _from_docx(data: bytes) -> Extraction:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        return Extraction.failed(f"Could not open the document ({type(exc).__name__}).")

    segments: list[Segment] = []
    heading: str | None = None
    body: list[str] = []

    def flush() -> None:
        if body:
            segments.append(Segment(text="\n\n".join(body), section=heading))

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        # A heading starts a new section, which becomes a chunk boundary and
        # the label a citation quotes.
        if paragraph.style is not None and str(paragraph.style.name).startswith("Heading"):
            flush()
            heading, body = text, []
        else:
            body.append(text)
    flush()

    for number, table in enumerate(document.tables, start=1):
        rows = [
            " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            for row in table.rows
        ]
        rows = [r for r in rows if r]
        if rows:
            segments.append(Segment(text="\n\n".join(rows), section=f"Table {number}"))

    return Extraction.of(segments)


def _read_page_images(page) -> str:  # noqa: ANN001 -- a pypdf PageObject
    if not ocr.available():
        return ""
    pieces = []
    try:
        images = list(page.images)
    except Exception:  # noqa: BLE001 -- some PDFs raise rather than list nothing
        return ""
    for image in images:
        try:
            recovered = ocr.read_image(image.data)
        except Exception:  # noqa: BLE001
            continue
        if recovered:
            pieces.append(recovered)
    return "\n\n".join(pieces)


def _from_pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        return Extraction.failed(f"Could not open the PDF ({type(exc).__name__}).")

    segments: list[Segment] = []
    scanned_pages = 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:  # noqa: BLE001 -- one broken page shouldn't lose the rest
            text = ""

        if len(text) >= MIN_PDF_PAGE_CHARS:
            segments.append(Segment(text=text, section=f"page {number}"))
            continue

        # Nothing selectable on this page: it is almost certainly a scan, so
        # read the picture of it.
        recovered = _read_page_images(page)
        if recovered:
            scanned_pages += 1
            segments.append(
                Segment(text=recovered, section=f"page {number}", source=ChunkSource.OCR)
            )
        elif text:
            segments.append(Segment(text=text, section=f"page {number}"))

    if not segments:
        if not ocr.available():
            return Extraction.empty(
                "No selectable text, and it could not be read as a scan because OCR is "
                "unavailable on this deployment."
            )
        return Extraction.empty("No text could be read from this PDF, including as a scan.")

    result = Extraction.of(segments)
    if scanned_pages and result.status == IndexStatus.INDEXED:
        result.detail = f"{scanned_pages} page(s) were read as scans using OCR."
    return result


def _from_image(data: bytes) -> Extraction:
    if not ocr.available():
        return Extraction.empty("OCR is unavailable on this deployment, so images aren't read.")
    text = ocr.read_image(data)
    if not text:
        return Extraction.empty("No legible text was found in this image.")
    return Extraction(
        segments=[Segment(text=text, source=ChunkSource.OCR)],
        status=IndexStatus.INDEXED,
        detail="Read from an image using OCR, so the text may be imperfect.",
    )


_HANDLERS = {
    ".txt": _from_text,
    ".md": _from_text,
    ".csv": _from_csv,
    ".xlsx": _from_xlsx,
    ".docx": _from_docx,
    ".pdf": _from_pdf,
}


def extract(filename: str, data: bytes) -> Extraction:
    """Text and structure for one file, or the reason there is none."""
    extension = Path(filename).suffix.lower()
    if legacy := LEGACY_OFFICE.get(extension):
        return Extraction.failed(
            f"{legacy} files can't be read. Re-save it in the modern format to make it "
            "searchable."
        )
    if extension in IMAGE_EXTENSIONS:
        return _from_image(data)

    handler = _HANDLERS.get(extension)
    if handler is None:
        return Extraction.failed(f"No reader for '{extension}' files.")
    try:
        return handler(data)
    except Exception as exc:  # noqa: BLE001 -- a malformed upload is not a server error
        return Extraction.failed(f"Could not read this file ({type(exc).__name__}).")
